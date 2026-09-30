import base64
import io
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent
A = {
    'Инвестиции в запуск, ₽': 2_400_000,
    'Средний чек, ₽': 390,
    'Чеков в день на старте': 70,
    'Чеков в день на плато': 190,
    'Месяцев до плато': 9,
    'Рабочих дней в месяце': 30,
    'Себестоимость, % от выручки': 0.31,
    'Аренда, ₽/мес': 180_000,
    'ФОТ, ₽/мес': 420_000,
    'Прочие постоянные, ₽/мес': 90_000,
    'Эквайринг и налог, % от выручки': 0.085,
    'Маркетинг на старте, ₽/мес': 120_000,
    'Маркетинг после плато, ₽/мес': 50_000,
}
N = 24
SEAS = [0.86, 0.9, 0.97, 1.02, 1.08, 1.12, 1.14, 1.1, 1.04, 0.98, 0.93, 1.02]

wb = Workbook()
ws = wb.active
ws.title = 'Допущения'
hdr = Font(bold=True, color='FFFFFF')
fill = PatternFill('solid', fgColor='1F3A5F')
inp = PatternFill('solid', fgColor='FFF4D6')
thin = Side(style='thin', color='D0D5DD')
ws['A1'] = 'Допущения модели'; ws['A1'].font = Font(bold=True, size=14)
for i, (k, v) in enumerate(A.items(), start=3):
    ws.cell(i, 1, k)
    c = ws.cell(i, 2, v); c.fill = inp; c.border = Border(thin, thin, thin, thin)
    c.number_format = '0.0%' if isinstance(v, float) else '# ##0'
ws.column_dimensions['A'].width = 38; ws.column_dimensions['B'].width = 16
ref = {k: f"Допущения!$B${i}" for i, k in enumerate(A, start=3)}

p = wb.create_sheet('P&L')
rows = ['Месяц', 'Чеков в день', 'Выручка', 'Себестоимость', 'Валовая прибыль', 'Аренда', 'ФОТ', 'Маркетинг',
        'Прочие постоянные', 'Эквайринг и налог', 'Операционная прибыль', 'Денежный поток накопленный']
for r, name in enumerate(rows, start=1):
    c = p.cell(r, 1, name); c.font = Font(bold=True) if r in (1, 3, 5, 11, 12) else Font()
p.column_dimensions['A'].width = 28
p.cell(14, 1, 'Сезонный коэффициент')
for m in range(1, N + 1):
    col = get_column_letter(m + 1)
    p.column_dimensions[col].width = 12
    p[f'{col}1'] = m; p[f'{col}1'].font = hdr; p[f'{col}1'].fill = fill; p[f'{col}1'].alignment = Alignment(horizontal='center')
    p[f'{col}2'] = (f"=IF({m}>={ref['Месяцев до плато']},{ref['Чеков в день на плато']},"
                    f"{ref['Чеков в день на старте']}+({ref['Чеков в день на плато']}-{ref['Чеков в день на старте']})*({m}-1)/({ref['Месяцев до плато']}-1))")
    p[f'{col}14'] = SEAS[(m - 1) % 12]; p[f'{col}14'].fill = inp; p[f'{col}14'].number_format = '0.00'
    p[f'{col}3'] = f"={col}2*{col}14*{ref['Средний чек, ₽']}*{ref['Рабочих дней в месяце']}"
    p[f'{col}4'] = f"=-{col}3*{ref['Себестоимость, % от выручки']}"
    p[f'{col}5'] = f"={col}3+{col}4"
    p[f'{col}6'] = f"=-{ref['Аренда, ₽/мес']}"
    p[f'{col}7'] = f"=-{ref['ФОТ, ₽/мес']}"
    p[f'{col}8'] = f"=-IF({m}>={ref['Месяцев до плато']},{ref['Маркетинг после плато, ₽/мес']},{ref['Маркетинг на старте, ₽/мес']})"
    p[f'{col}9'] = f"=-{ref['Прочие постоянные, ₽/мес']}"
    p[f'{col}10'] = f"=-{col}3*{ref['Эквайринг и налог, % от выручки']}"
    p[f'{col}11'] = f"=SUM({col}5:{col}10)"
    prev = get_column_letter(m)
    p[f'{col}12'] = f"=-{ref['Инвестиции в запуск, ₽']}+{col}11" if m == 1 else f"={prev}12+{col}11"
    for r in range(2, 13):
        p[f'{col}{r}'].number_format = '# ##0;[Red]-# ##0'
s = wb.create_sheet('Итоги')
last = get_column_letter(N + 1)
s['A1'] = 'Ключевые показатели'; s['A1'].font = Font(bold=True, size=14)
kpi = [('Выручка за 2 года, ₽', f"=SUM('P&L'!B3:{last}3)"),
       ('Операционная прибыль за 2 года, ₽', f"=SUM('P&L'!B11:{last}11)"),
       ('Первый прибыльный месяц', f"=IFERROR(MATCH(TRUE,INDEX('P&L'!B11:{last}11>0,0),0),\"нет\")"),
       ('Месяц окупаемости', f"=IFERROR(MATCH(TRUE,INDEX('P&L'!B12:{last}12>=0,0),0),\"> 24\")"),
       ('Точка безубыточности, чеков в день',
        f"=({ref['Аренда, ₽/мес']}+{ref['ФОТ, ₽/мес']}+{ref['Прочие постоянные, ₽/мес']}+{ref['Маркетинг после плато, ₽/мес']})/"
        f"({ref['Средний чек, ₽']}*{ref['Рабочих дней в месяце']}*(1-{ref['Себестоимость, % от выручки']}-{ref['Эквайринг и налог, % от выручки']}))")]
for i, (k, f) in enumerate(kpi, start=3):
    s.cell(i, 1, k); c = s.cell(i, 2, f); c.number_format = '# ##0'
s.column_dimensions['A'].width = 40; s.column_dimensions['B'].width = 16
ch = LineChart(); ch.title = 'Накопленный денежный поток'; ch.height = 8; ch.width = 18
ch.add_data(Reference(p, min_col=2, max_col=N + 1, min_row=12), from_rows=True, titles_from_data=False)
s.add_chart(ch, 'D3')
bc = BarChart(); bc.title = 'Операционная прибыль по месяцам'; bc.height = 8; bc.width = 18
bc.add_data(Reference(p, min_col=2, max_col=N + 1, min_row=11), from_rows=True, titles_from_data=False)
s.add_chart(bc, 'D20')
wb.save(OUT / 'finmodel_coffee.xlsx')

v = A
checks, rev, op, cum = [], [], [], []
c0 = -v['Инвестиции в запуск, ₽']
for m in range(1, N + 1):
    ck = v['Чеков в день на плато'] if m >= v['Месяцев до плато'] else v['Чеков в день на старте'] + (v['Чеков в день на плато'] - v['Чеков в день на старте']) * (m - 1) / (v['Месяцев до плато'] - 1)
    r = ck * SEAS[(m - 1) % 12] * v['Средний чек, ₽'] * v['Рабочих дней в месяце']
    mk = v['Маркетинг после плато, ₽/мес'] if m >= v['Месяцев до плато'] else v['Маркетинг на старте, ₽/мес']
    o = r * (1 - v['Себестоимость, % от выручки'] - v['Эквайринг и налог, % от выручки']) - v['Аренда, ₽/мес'] - v['ФОТ, ₽/мес'] - mk - v['Прочие постоянные, ₽/мес']
    c0 += o
    checks.append(ck); rev.append(r); op.append(o); cum.append(c0)
first_profit = next(i + 1 for i, x in enumerate(op) if x > 0)
payback = next((i + 1 for i, x in enumerate(cum) if x >= 0), None)
be = (v['Аренда, ₽/мес'] + v['ФОТ, ₽/мес'] + v['Прочие постоянные, ₽/мес'] + v['Маркетинг после плато, ₽/мес']) / (v['Средний чек, ₽'] * v['Рабочих дней в месяце'] * (1 - v['Себестоимость, % от выручки'] - v['Эквайринг и налог, % от выручки']))

plt.rcParams.update({'font.family': 'Arial', 'axes.spines.top': False, 'axes.spines.right': False, 'axes.edgecolor': '#98a2b3'})
fig, ax = plt.subplots(figsize=(9, 3.8))
ms = list(range(1, N + 1))
ax.bar(ms, [x / 1e3 for x in op], color=['#e5484d' if x < 0 else '#2f6df6' for x in op], width=0.7)
ax2 = ax.twinx(); ax2.plot(ms, [x / 1e6 for x in cum], color='#12b76a', lw=2.5); ax2.axhline(0, color='#12b76a', lw=0.8, ls='--')
a0, a1 = ax.get_ylim(); rt = -a0 / (a1 - a0); b1 = max(cum) / 1e6 * 1.1; b0 = -rt / (1 - rt) * b1
if b0 > min(cum) / 1e6 * 1.1:
    b0 = min(cum) / 1e6 * 1.1; b1 = -b0 * (1 - rt) / rt
ax2.set_ylim(b0, b1)
ax.set_xlabel('месяц'); ax.set_ylabel('прибыль, тыс ₽'); ax2.set_ylabel('накоплено, млн ₽'); ax2.spines['right'].set_visible(True)
if payback: ax2.annotate(f'окупаемость: {payback}-й месяц', (payback, 0), xytext=(1.2, max(cum) / 1e6 * 0.95), arrowprops=dict(arrowstyle='->', color='#12b76a'), color='#0b7a4b', fontsize=11)
buf = io.BytesIO(); fig.savefig(buf, format='png', dpi=170, bbox_inches='tight'); plt.close(fig)
g = base64.b64encode(buf.getvalue()).decode()
fmt = lambda x: f'{x:,.0f}'.replace(',', ' ')
tbl = ''.join(f'<tr><td>{m}</td><td>{checks[m-1]:.0f}</td><td>{fmt(rev[m-1])}</td><td class="{"n" if op[m-1] < 0 else "p"}">{fmt(op[m-1])}</td><td class="{"n" if cum[m-1] < 0 else "p"}">{fmt(cum[m-1])}</td></tr>' for m in (1, 3, 6, 9, 12, 18, 24))
html = f'''<!doctype html><html lang="ru"><head><meta charset="utf-8"><link href="https://fonts.googleapis.com/css2?family=Manrope:wght@500;700;800&display=swap" rel="stylesheet">
<style>*{{margin:0;box-sizing:border-box}}body{{width:1600px;background:#f3f5f8;font-family:Manrope;color:#101828;padding:60px}}
h1{{font:800 46px/1.1 Manrope;letter-spacing:-1px}}.sub{{color:#475467;font-size:19px;margin:12px 0 34px}}
.k{{display:grid;grid-template-columns:repeat(4,1fr);gap:16px;margin-bottom:22px}}.k div{{background:#fff;border-radius:16px;padding:22px}}.k b{{font:800 34px Manrope;display:block}}.k span{{color:#475467;font-size:15px}}
.g{{display:grid;grid-template-columns:1.35fr 1fr;gap:22px}}.c{{background:#fff;border-radius:16px;padding:20px}}.c img{{width:100%}}
table{{width:100%;border-collapse:collapse;font-size:16px}}th{{text-align:right;color:#475467;font-weight:700;padding:10px 8px;border-bottom:2px solid #e4e7ec}}td{{text-align:right;padding:10px 8px;border-bottom:1px solid #f2f4f7}}th:first-child,td:first-child{{text-align:left}}
.n{{color:#d92d20}}.p{{color:#067647}}.note{{margin-top:22px;display:grid;grid-template-columns:repeat(3,1fr);gap:16px}}.note div{{background:#101828;color:#e4e7ec;border-radius:16px;padding:22px;font-size:16px;line-height:1.5}}.note b{{color:#fff;display:block;margin-bottom:6px}}</style></head><body>
<h1>Финансовая модель кофейни на 24 месяца</h1><div class="sub">Excel с формулами: допущения → P&amp;L по месяцам → денежный поток → итоги. Меняете средний чек или аренду — пересчитывается всё.</div>
<div class="k"><div><b>{fmt(v['Инвестиции в запуск, ₽'])} ₽</b><span>вложения в запуск</span></div><div><b>{first_profit}-й</b><span>первый прибыльный месяц</span></div><div><b>{payback or '> 24'}-й</b><span>месяц окупаемости</span></div><div><b>{be:.0f}</b><span>чеков в день — безубыточность</span></div></div>
<div class="g"><div class="c"><img src="data:image/png;base64,{g}"></div><div class="c"><table><tr><th>Месяц</th><th>Чеков/день</th><th>Выручка</th><th>Прибыль</th><th>Накоплено</th></tr>{tbl}</table></div></div>
<div class="note"><div><b>Сценарии</b>Лист допущений с жёлтыми ячейками: меняете цифру — видите, как сдвигается окупаемость.</div><div><b>Для банка и инвестора</b>P&amp;L, денежный поток и точка безубыточности в одном файле.</div><div><b>Чистые формулы</b>Без макросов и скрытых листов, открывается в Excel и Google Таблицах.</div></div>
</body></html>'''
(OUT / 'show.html').write_text(html, encoding='utf-8')
print('ok', first_profit, payback, round(be))
